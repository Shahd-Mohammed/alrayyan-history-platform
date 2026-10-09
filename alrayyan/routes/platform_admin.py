import csv
import json
import os
import re
import smtplib
from email.message import EmailMessage
from datetime import datetime, timezone
from pathlib import Path
from uuid import uuid4
from io import BytesIO, StringIO

from flask import Blueprint, abort, current_app, flash, redirect, render_template, request, url_for
from flask_login import current_user, login_required

from alrayyan.extensions import db
from alrayyan.forms import ClassroomForm, ConceptMapForm, ConceptMapUploadForm, CurriculumUploadForm, HistoricalCharacterForm, InvitationForm, LearningResourceForm, PlatformSettingsForm
from alrayyan.models import CharacterAttempt, Classroom, ConceptMap, ConceptMapEdge, ConceptMapNode, ConceptMastery, ContentChunk, Curriculum, HistoricalCharacter, HonorBoardEntry, LearningResource, Lesson, PlatformSettings, SourceDocument, StudentInvitation, TutorConversation, TutorMessage, LearningPlanItem, Unit, User, Worksheet, WorksheetAttempt, ChallengeSession, XPTransaction
from alrayyan.services.document_reader import calculate_checksum, extract_docx, extract_pdf
from alrayyan.services.text_processing import chunk_text, create_text_hash
from alrayyan.services.curriculum_processing import process_source_document
from alrayyan.services.learning_content_ai import (
    generate_character_drafts,
    generate_concept_map_draft,
)
from alrayyan.services.embeddings import generate_embeddings, get_embedding_settings

platform_admin_bp = Blueprint("platform_admin", __name__, url_prefix="/teacher-dashboard/platform")


@platform_admin_bp.before_request
@login_required
def teacher_only():
    if current_user.role not in {"teacher", "admin"}:
        abort(403)


def teacher_curricula():
    query = Curriculum.query
    if current_user.role != "admin":
        query = query.filter((Curriculum.created_by_id == current_user.id) | (Curriculum.created_by_id.is_(None)))
    return query


def teacher_lessons():
    # Lesson selectors include every curriculum that the teacher has
    # explicitly activated. Activation is independent per curriculum.
    # Inactive curricula must never leak into teacher-facing dropdowns.
    # Ignore legacy placeholder lessons left by older curriculum parsing.
    placeholder_titles = {
        "محتوى المنهج",
        "الدرس احتياطي",
        "الوحدة — الدرس احتياطي",
    }
    query = Lesson.query.join(Unit).join(Curriculum).filter(
        Curriculum.is_active.is_(True),
        Curriculum.processing_status == "ready",
        Lesson.title.notin_(placeholder_titles),
    )

    # Admins manage shared platform curricula, including curricula uploaded
    # by another teacher. Non-admin teachers only see their own or shared
    # (ownerless) curricula, matching teacher_curricula().
    if current_user.role != "admin":
        query = query.filter(
            (Curriculum.created_by_id == current_user.id)
            | (Curriculum.created_by_id.is_(None))
        )

    return query.order_by(
        Curriculum.subject,
        Curriculum.grade,
        Unit.order_index,
        Lesson.order_index,
    ).all()


def lesson_choices():
    return [(lesson.id, f"{lesson.unit.title} — {lesson.title}") for lesson in teacher_lessons()]


def curriculum_ready_for_activation(curriculum):
    if curriculum.processing_status != "ready" or not curriculum.sources:
        return False
    chunks = ContentChunk.query.join(SourceDocument).filter(
        SourceDocument.curriculum_id == curriculum.id,
    ).all()
    return bool(chunks) and all(
        chunk.embedding and chunk.embedding_model and chunk.embedding_dimensions
        for chunk in chunks
    )


def normalize_active_curricula():
    """Compatibility hook: active curricula are controlled manually."""
    return {
        (curriculum.subject, curriculum.grade): curriculum
        for curriculum in Curriculum.query.filter_by(is_active=True).all()
    }


def activate_curriculum(curriculum):
    if not curriculum_ready_for_activation(curriculum):
        raise ValueError("المنهج لم يكتمل تجهيزه وفهرسته: يجب أن تكون كل المقاطع مستخرجة ومفهرسة بالـembeddings قبل التفعيل.")

    # Activation is an independent manual choice. Enabling one curriculum
    # must never deactivate another curriculum.
    curriculum.is_active = True
    curriculum.processing_status = "ready"
    curriculum.activated_at = datetime.now(timezone.utc).replace(tzinfo=None)

    for source in curriculum.sources:
        source.is_active = True


def owned_lesson_or_404(lesson_id):
    lesson = db.session.get(Lesson, lesson_id)
    if lesson is None or lesson.id not in {item.id for item in teacher_lessons()}:
        abort(404)
    return lesson


def save_upload(upload, folder_name, allowed):
    original = (upload.filename or "").strip()
    extension = Path(original).suffix.lower()
    if extension.lstrip(".") not in allowed:
        raise ValueError("صيغة الملف غير مسموحة")
    folder = Path(current_app.config["UPLOAD_FOLDER"]) / folder_name
    folder.mkdir(parents=True, exist_ok=True)
    stored_name = f"{uuid4().hex}{extension}"
    path = folder / stored_name
    upload.save(path)
    return original, path, str(Path(folder_name) / stored_name)


@platform_admin_bp.route("/curricula", methods=["GET", "POST"])
def curricula():
    form = CurriculumUploadForm()
    if form.validate_on_submit():
        stored_path = None
        try:
            stored_path = None
            original, stored_path, relative = save_upload(form.document.data, "curricula", {"pdf", "docx"})
            curriculum_query = teacher_curricula().filter_by(
                subject=form.subject.data.strip(),
                grade=form.grade.data.strip(),
                semester=form.semester.data,
                academic_year=form.academic_year.data.strip(),
            ).order_by(Curriculum.created_at.desc())
            curriculum = curriculum_query.first()
            if curriculum is None or curriculum.is_active:
                versions = [str(item.version) for item in curriculum_query.all()]
                version_numbers = []
                for version in versions:
                    try:
                        version_numbers.append(float(version))
                    except (TypeError, ValueError):
                        continue
                next_version = f"{(max(version_numbers) + 1.0) if version_numbers else 1.0:.1f}"
                curriculum = Curriculum(
                    created_by_id=current_user.id,
                    name=(form.curriculum_name.data or "").strip() or f"{form.subject.data.strip()} — {form.grade.data.strip()} — الفصل {form.semester.data} — {form.academic_year.data.strip()}",
                    subject=form.subject.data.strip(),
                    grade=form.grade.data.strip(),
                    semester=form.semester.data,
                    academic_year=form.academic_year.data.strip(),
                    version=next_version,
                    is_active=False,
                    processing_status="processing",
                )
                db.session.add(curriculum)
                db.session.flush()
            else:
                curriculum.processing_status = "processing"
                curriculum.processing_error = None
                curriculum.is_active = False
                for source in curriculum.sources:
                    source.is_active = False

            checksum = calculate_checksum(stored_path)
            if SourceDocument.query.filter_by(curriculum_id=curriculum.id, checksum=checksum).first():
                raise ValueError("هذا الملف موجود مسبقًا داخل المنهج نفسه")

            extracted = (
                extract_pdf(stored_path)
                if stored_path.suffix.lower() == ".pdf"
                else extract_docx(stored_path)
            )

            source_priority = {
                "official_book": 1,
                "supporting_book": 2,
                "review_notes": 3,
            }[form.source_type.data]

            source = SourceDocument(
                curriculum_id=curriculum.id,
                lesson_id=None,
                title=(form.source_title.data or "").strip() or Path(original).stem,
                source_type=form.source_type.data,
                original_filename=original,
                stored_path=relative,
                page_count=extracted.get("page_count"),
                checksum=checksum,
                academic_year=curriculum.academic_year,
                priority=source_priority,
                is_primary=form.source_type.data == "official_book",
                is_active=True,
            )
            db.session.add(source)
            db.session.flush()

            extracted_count = process_source_document(
                source,
                fallback_unit_title=(form.unit_title.data or "").strip(),
                fallback_lesson_title=(form.lesson_title.data or "").strip(),
            )
            # Uploading and indexing a curriculum does not activate it.
            # Activation remains an explicit manual choice from the curriculum
            # management screen.
            curriculum.is_active = False
            for item in curriculum.sources:
                item.is_active = False
            curriculum.processing_status = "ready"
            db.session.commit()
            flash(
                f"تم تحليل وفهرسة {extracted_count} مقطعًا. المنهج جاهز ويمكنك تفعيله يدويًا من قائمة المناهج.",
                "success",
            )
            return redirect(url_for("platform_admin.curricula"))
        except Exception as error:
            db.session.rollback()
            if stored_path and stored_path.exists():
                stored_path.unlink()
            current_app.logger.exception("Curriculum upload failed")
            flash(f"تعذر تجهيز المصدر: {error}", "error")

    normalize_active_curricula()

    return render_template(
        "platform/curricula.html",
        form=form,
        curricula=teacher_curricula().order_by(Curriculum.created_at.desc()).all(),
    )


@platform_admin_bp.post("/curricula/<int:curriculum_id>/reindex")
def reindex_curriculum(curriculum_id):
    curriculum = teacher_curricula().filter_by(id=curriculum_id).first_or_404()
    if not curriculum.sources:
        flash("لا يوجد ملف مصدر لإعادة تجهيزه داخل هذا المنهج.", "error")
        return redirect(url_for("platform_admin.curricula"))

    try:
        curriculum.processing_status = "processing"
        curriculum.processing_error = None
        total_chunks = 0
        for source in curriculum.sources:
            total_chunks += process_source_document(source)

        curriculum.processing_status = "ready"
        db.session.commit()
        flash(
            f"تمت إعادة تجهيز وفهرسة {total_chunks} مقطعًا. يمكنك الآن تفعيل المنهج.",
            "success",
        )
    except Exception as error:
        db.session.rollback()
        current_app.logger.exception("Curriculum reindex failed")
        flash(f"تعذر إعادة تجهيز المنهج: {error}", "error")

    return redirect(url_for("platform_admin.curricula"))


@platform_admin_bp.post("/curricula/<int:curriculum_id>/delete")
def delete_curriculum(curriculum_id):
    curriculum = teacher_curricula().filter_by(id=curriculum_id).first_or_404()
    relative_paths = [source.stored_path for source in curriculum.sources if source.stored_path]
    was_active = curriculum.is_active

    try:
        db.session.delete(curriculum)
        db.session.commit()
    except Exception as error:
        db.session.rollback()
        current_app.logger.exception("Curriculum deletion failed")
        flash(f"تعذر حذف المنهج: {error}", "error")
        return redirect(url_for("platform_admin.curricula"))

    upload_root = Path(current_app.config["UPLOAD_FOLDER"]).resolve()
    for relative_path in relative_paths:
        target = (upload_root / relative_path).resolve()
        if target.is_relative_to(upload_root) and target.is_file():
            try:
                target.unlink()
            except OSError:
                current_app.logger.warning("Could not remove curriculum file: %s", target)

    flash(
        "تم حذف المنهج ومصادره وملفاته نهائيًا."
        if was_active
        else "تم حذف المنهج ومصادره وملفاته نهائيًا.",
        "success",
    )
    return redirect(url_for("platform_admin.curricula"))


@platform_admin_bp.post("/curricula/<int:curriculum_id>/toggle")
def toggle_curriculum(curriculum_id):
    curriculum = teacher_curricula().filter_by(id=curriculum_id).first_or_404()
    new_state = not curriculum.is_active
    if new_state:
        try:
            activate_curriculum(curriculum)
            db.session.commit()
        except Exception as error:
            db.session.rollback()
            flash(str(error), "error")
            return redirect(url_for("platform_admin.curricula"))
    else:
        curriculum.is_active = False
        for source in curriculum.sources:
            source.is_active = False

        # Deactivation is also independent. Do not automatically
        # activate another curriculum when this one is switched off.
        db.session.commit()
    flash("تم تحديث حالة المنهج ومصادره", "success")
    return redirect(url_for("platform_admin.curricula"))


@platform_admin_bp.route("/classrooms", methods=["GET", "POST"])
def classrooms():
    form = ClassroomForm()
    if form.validate_on_submit():
        db.session.add(Classroom(teacher_id=current_user.id, name=form.name.data.strip(), grade=form.grade.data.strip(), academic_year=form.academic_year.data.strip(), whatsapp_url=(form.whatsapp_url.data or "").strip() or None))
        db.session.commit(); flash("تم إنشاء الصف", "success")
        return redirect(url_for("platform_admin.classrooms"))
    query = Classroom.query if current_user.role == "admin" else Classroom.query.filter_by(teacher_id=current_user.id)
    return render_template("platform/classrooms.html", form=form, classrooms=query.order_by(Classroom.created_at.desc()).all())


@platform_admin_bp.post("/classrooms/<int:classroom_id>/delete")
def delete_classroom(classroom_id):
    classroom = Classroom.query.filter_by(id=classroom_id, teacher_id=current_user.id).first_or_404()
    try:
        # Removing a class only removes enrollments; student accounts and worksheet results stay intact.
        db.session.delete(classroom)
        db.session.commit()
        flash("تم حذف الصف وإزالة ارتباط الطالبات به، مع الحفاظ على حساباتهن ونتائجهن.", "success")
    except Exception:
        db.session.rollback()
        current_app.logger.exception("Classroom deletion failed")
        flash("تعذر حذف الصف. لم يتم تغيير البيانات.", "error")
    return redirect(url_for("platform_admin.classrooms"))


def invitation_mail_settings():
    return all(os.getenv(key) for key in ("SMTP_HOST", "SMTP_USERNAME", "SMTP_PASSWORD", "SMTP_FROM_EMAIL"))


def send_student_invitation(email, invitation_url, teacher_name):
    host, username, password, sender = (os.getenv(key) for key in ("SMTP_HOST", "SMTP_USERNAME", "SMTP_PASSWORD", "SMTP_FROM_EMAIL"))
    if not all((host, username, password, sender)):
        raise RuntimeError("إعدادات البريد غير مكتملة")
    message = EmailMessage()
    message["Subject"] = "دعوة للانضمام إلى منصة الريان للدراسات التاريخية"
    message["From"] = sender
    message["To"] = email
    message.set_content(f"مرحبًا،\n\nدعتك المعلمة {teacher_name} للانضمام إلى منصة الريان للدراسات التاريخية.\nافتحي الرابط التالي لإنشاء حسابك: {invitation_url}\n\nإذا لم تتوقعي هذه الدعوة، يمكنك تجاهل هذه الرسالة.")
    with smtplib.SMTP(host, int(os.getenv("SMTP_PORT", "587")), timeout=20) as server:
        if os.getenv("SMTP_USE_TLS", "1").lower() not in {"0", "false", "no"}:
            server.starttls()
        server.login(username, password)
        server.send_message(message)


def extract_invitation_emails(upload):
    extension = Path(upload.filename or "").suffix.lower()
    if extension not in {".xlsx", ".csv"}:
        raise ValueError("ارفعي ملف Excel بصيغة ‎.xlsx أو ملف ‎.csv.")
    raw = upload.read()
    values = []
    if extension == ".csv":
        try:
            content = raw.decode("utf-8-sig")
        except UnicodeDecodeError as error:
            raise ValueError("تعذر قراءة ملف CSV؛ احفظيه بترميز UTF-8.") from error
        for row in csv.reader(StringIO(content)):
            values.extend(str(cell).strip() for cell in row if cell is not None)
    else:
        try:
            from openpyxl import load_workbook
            workbook = load_workbook(BytesIO(raw), read_only=True, data_only=True)
            for sheet in workbook.worksheets:
                for row in sheet.iter_rows(values_only=True):
                    values.extend(str(cell).strip() for cell in row if cell is not None)
            workbook.close()
        except Exception as error:
            raise ValueError("تعذر قراءة ملف Excel. تأكدي أنه ملف ‎.xlsx سليم.") from error
    pattern = re.compile(r"[^@\s]+@[^@\s]+\.[^@\s]+")
    emails = []
    malformed = []
    for value in values:
        candidates = re.split(r"[\s,;]+", value.lower())
        for candidate in candidates:
            candidate = candidate.strip().strip("()<>[]{}.,")
            if not candidate:
                continue
            if "@" in candidate:
                if pattern.fullmatch(candidate):
                    if candidate not in emails:
                        emails.append(candidate)
                else:
                    malformed.append(candidate)
    if malformed:
        raise ValueError("يوجد عنوان بريد غير صالح في الملف: " + "، ".join(malformed[:5]))
    if not emails:
        raise ValueError("لم أجد عناوين بريد إلكتروني صالحة داخل الملف.")
    return emails


@platform_admin_bp.route("/invitations", methods=["GET", "POST"])
def invitations():
    form = InvitationForm()
    classes = Classroom.query.filter_by(teacher_id=current_user.id, is_active=True).all()
    form.classroom_id.choices = [(0, "بدون صف محدد")] + [(row.id, row.name) for row in classes]
    invitation_urls = []
    upload = request.files.get("emails_file")
    is_file_upload = request.method == "POST" and upload is not None and bool(upload.filename)
    if is_file_upload:
        try:
            emails = extract_invitation_emails(upload)
            classroom_id = request.form.get("classroom_id", default=0, type=int) or None
            if classroom_id and not any(row.id == classroom_id for row in classes):
                abort(403)
            valid_days = request.form.get("valid_days", default=7, type=int)
            if not 1 <= valid_days <= 30:
                raise ValueError("مدة صلاحية الدعوة يجب أن تكون بين يوم و30 يومًا.")
            invitation_urls = create_invitation_batch(emails, classroom_id, valid_days)
        except ValueError as error:
            flash(str(error), "error")
        except Exception as error:
            db.session.rollback()
            current_app.logger.exception("Bulk invitation upload failed")
            flash(f"تعذر معالجة الملف: {error}", "error")
    elif form.validate_on_submit():
        emails = []
        for email in re.split(r"[\s,;]+", form.emails.data or ""):
            normalized = email.strip().lower()
            if normalized and normalized not in emails:
                emails.append(normalized)
        invalid = [email for email in emails if not re.fullmatch(r"[^@\s]+@[^@\s]+\.[^@\s]+", email)]
        if invalid:
            form.emails.errors.append("بعض العناوين غير صالحة: " + "، ".join(invalid[:5]))
        else:
            invitation_urls = create_invitation_batch(emails, form.classroom_id.data or None, form.valid_days.data)
    query = StudentInvitation.query if current_user.role == "admin" else StudentInvitation.query.filter_by(teacher_id=current_user.id)
    return render_template("platform/invitations.html", form=form, invitations=query.order_by(StudentInvitation.created_at.desc()).limit(200).all(), invitation_urls=invitation_urls, classrooms=classes)


def create_invitation_batch(emails, classroom_id, valid_days):
    existing = {row.email for row in User.query.filter(User.email.in_(emails)).all()} if emails else set()
    pending = StudentInvitation.query.filter(
        StudentInvitation.teacher_id == current_user.id,
        StudentInvitation.email.in_(emails),
        StudentInvitation.accepted_at.is_(None),
        StudentInvitation.expires_at > datetime.utcnow(),
    ).all() if emails else []
    existing_invites = {row.email for row in pending}
    invitation_urls = []
    for email in emails:
        if email in existing or email in existing_invites:
            continue
        invitation, token = StudentInvitation.create(current_user.id, email, classroom_id, valid_days)
        db.session.add(invitation)
        invitation_urls.append({"email": email, "url": url_for("auth.register_invitation", token=token, _external=True)})
    db.session.commit()
    if not invitation_urls:
        flash(f"لم تُنشأ دعوات جديدة؛ تم تجاوز {len(existing)} حسابًا موجودًا و{len(existing_invites)} دعوة سارية.", "warning")
        return []
    if invitation_mail_settings():
        sent, failed = 0, []
        for item in invitation_urls:
            try:
                send_student_invitation(item["email"], item["url"], current_user.full_name or "معلمة منصة الريان")
                sent += 1
            except Exception:
                current_app.logger.exception("Could not email student invitation")
                failed.append(item["email"])
        if failed:
            flash(f"تم إنشاء {len(invitation_urls)} دعوة. أُرسلت {sent} رسالة، وتعذر إرسال {len(failed)}؛ الروابط متاحة للنسخ أدناه.", "warning")
        else:
            flash(f"تم إنشاء وإرسال {sent} دعوة بالبريد الإلكتروني بنجاح.", "success")
    else:
        flash(f"تم إنشاء {len(invitation_urls)} دعوة، لكن إرسال البريد غير مفعّل بعد. أضيفي إعدادات SMTP في Render ليرسل النظام الرسائل تلقائيًا؛ الروابط متاحة للنسخ أدناه.", "warning")
    return invitation_urls


@platform_admin_bp.route("/resources", methods=["GET", "POST"])
def resources():
    form = LearningResourceForm(); form.lesson_id.choices = lesson_choices()
    if form.validate_on_submit():
        try:
            owned_lesson_or_404(form.lesson_id.data)
            original, path, relative = save_upload(form.document.data, "learning_resources", {"pdf", "doc", "docx", "png", "jpg", "jpeg"})
            resource = LearningResource(lesson_id=form.lesson_id.data, created_by_id=current_user.id, title=form.title.data.strip(), resource_type=form.resource_type.data, description=(form.description.data or "").strip() or None, original_filename=original, stored_path=relative, mime_type=form.document.data.mimetype, allow_download=form.allow_download.data, publication_status=form.publication_status.data)
            db.session.add(resource); db.session.commit(); flash("تم حفظ المحتوى التعليمي", "success")
            return redirect(url_for("platform_admin.resources"))
        except Exception as error:
            db.session.rollback(); flash(f"تعذر رفع الملف: {error}", "error")
    query = LearningResource.query if current_user.role == "admin" else LearningResource.query.filter_by(created_by_id=current_user.id)
    return render_template("platform/resources.html", form=form, resources=query.order_by(LearningResource.created_at.desc()).all())


def owned_content_or_404(model, row_id):
    row = db.session.get(model, row_id)
    if row is None:
        abort(404)
    if current_user.role != "admin" and row.created_by_id != current_user.id:
        abort(403)
    return row


@platform_admin_bp.post("/resources/<int:row_id>/toggle")
def toggle_resource(row_id):
    row = owned_content_or_404(LearningResource, row_id)
    row.publication_status = "draft" if row.publication_status == "published" else "published"
    db.session.commit()
    flash("تم تحديث ظهور الملف للطلاب.", "success")
    return redirect(url_for("platform_admin.resources"))


@platform_admin_bp.post("/resources/<int:row_id>/delete")
def delete_resource(row_id):
    row = owned_content_or_404(LearningResource, row_id)
    relative_path = row.stored_path
    db.session.delete(row)
    db.session.commit()
    if relative_path:
        target = (Path(current_app.config["UPLOAD_FOLDER"]) / relative_path).resolve()
        root = Path(current_app.config["UPLOAD_FOLDER"]).resolve()
        if target.is_relative_to(root) and target.is_file():
            target.unlink()
    flash("تم حذف الملف نهائيًا.", "success")
    return redirect(url_for("platform_admin.resources"))


@platform_admin_bp.route("/concept-maps", methods=["GET", "POST"])
def concept_maps():
    form = ConceptMapForm(); form.lesson_id.choices = lesson_choices()
    upload_form = ConceptMapUploadForm(prefix="upload"); upload_form.lesson_id.choices = lesson_choices()
    if form.validate_on_submit():
        try:
            graph = json.loads(form.nodes_json.data or "{}")
        except ValueError:
            graph = {}
        nodes_data = graph.get("nodes", []) if isinstance(graph, dict) else []
        edges_data = graph.get("edges", []) if isinstance(graph, dict) else []
        if len(nodes_data) < 2:
            form.nodes_json.errors.append("أضيفي مفهومين على الأقل داخل المصمم")
        else:
            owned_lesson_or_404(form.lesson_id.data)
            nodes_data = nodes_data[:80]
            edges_data = edges_data[:160]
            concept_map = ConceptMap(lesson_id=form.lesson_id.data, created_by_id=current_user.id, title=form.title.data.strip(), map_type="concept", creation_method="manual", publication_status=form.publication_status.data)
            db.session.add(concept_map); db.session.flush()
            nodes = {}
            for index, item in enumerate(nodes_data):
                label = str(item.get("label", "")).strip()[:250]
                if not label: continue
                try:
                    position_x = min(max(float(item.get("x", (index % 3) * 260)), 0), 2000)
                    position_y = min(max(float(item.get("y", (index // 3) * 150)), 0), 2000)
                except (TypeError, ValueError):
                    position_x, position_y = (index % 3) * 260, (index // 3) * 150
                node = ConceptMapNode(concept_map_id=concept_map.id, label=label, description=str(item.get("description", "")).strip()[:2000] or None, position_x=position_x, position_y=position_y)
                db.session.add(node); db.session.flush(); nodes[str(item.get("key", index))] = node
            for item in edges_data:
                source, target = nodes.get(str(item.get("source"))), nodes.get(str(item.get("target")))
                if source and target and source.id != target.id:
                    db.session.add(ConceptMapEdge(concept_map_id=concept_map.id, source_node_id=source.id, target_node_id=target.id, label=str(item.get("label", ""))[:150] or None))
            db.session.commit(); flash("تم إنشاء الخريطة المفاهيمية", "success")
            return redirect(url_for("platform_admin.concept_maps"))
    query = ConceptMap.query if current_user.role == "admin" else ConceptMap.query.filter_by(created_by_id=current_user.id)
    return render_template("platform/concept_maps.html", form=form, upload_form=upload_form, maps=query.order_by(ConceptMap.created_at.desc()).all())


@platform_admin_bp.post("/concept-maps/upload")
def upload_concept_map():
    form = ConceptMapUploadForm(prefix="upload"); form.lesson_id.choices = lesson_choices()
    if not form.validate_on_submit():
        flash("راجعي بيانات خريطة الرفع وصيغة الملف.", "error")
        return redirect(url_for("platform_admin.concept_maps"))
    owned_lesson_or_404(form.lesson_id.data)
    original, path, relative = save_upload(form.document.data, "concept_maps", {"pdf", "png", "jpg", "jpeg", "webp"})
    row = ConceptMap(lesson_id=form.lesson_id.data, created_by_id=current_user.id, title=form.title.data.strip(), map_type="uploaded", creation_method="upload", publication_status=form.publication_status.data, original_filename=original, stored_path=relative, mime_type=form.document.data.mimetype)
    db.session.add(row); db.session.commit()
    flash("تم رفع الخريطة الجاهزة.", "success")
    return redirect(url_for("platform_admin.concept_maps"))


def layout_concept_tree(node_by_key, edge_rows):
    """Store a tree-shaped layout so generated maps open as branches, not a grid."""
    children = {key: [] for key in node_by_key}
    targeted = set()
    for edge in edge_rows:
        source_key = str(edge.get("source"))
        target_key = str(edge.get("target"))
        if source_key in node_by_key and target_key in node_by_key and source_key != target_key:
            if target_key not in targeted:
                children[source_key].append(target_key)
                targeted.add(target_key)

    leaf_index = 0
    visited = set()

    def place(key, depth):
        nonlocal leaf_index
        if key in visited:
            return leaf_index * 155
        visited.add(key)
        node = node_by_key[key]
        node.position_x = depth * 300
        child_keys = [child for child in children.get(key, []) if child not in visited]
        if not child_keys:
            node.position_y = leaf_index * 155
            leaf_index += 1
            return node.position_y
        child_positions = [place(child, depth + 1) for child in child_keys]
        node.position_y = (child_positions[0] + child_positions[-1]) / 2
        return node.position_y

    roots = [key for key in node_by_key if key not in targeted]
    for key in roots:
        place(key, 0)
    for key in node_by_key:
        if key not in visited:
            place(key, 0)


@platform_admin_bp.post("/concept-maps/generate")
def generate_concept_map():
    lesson = owned_lesson_or_404(request.form.get("lesson_id", type=int))
    try:
        payload = generate_concept_map_draft(lesson)
        row = ConceptMap(
            lesson_id=lesson.id,
            created_by_id=current_user.id,
            title=payload.get("title") or f"خريطة {lesson.title}",
            map_type="concept",
            creation_method="ai",
            publication_status="draft",
        )
        db.session.add(row)
        db.session.flush()
        node_by_key = {}
        for item in payload["nodes"]:
            label = str(item.get("label", "")).strip()[:250]
            if not label:
                continue
            node = ConceptMapNode(
                concept_map_id=row.id,
                label=label,
                description=str(item.get("description", "")).strip() or None,
                position_x=0,
                position_y=0,
            )
            db.session.add(node)
            db.session.flush()
            node_key = str(item.get("key") or f"n{len(node_by_key) + 1}")
            while node_key in node_by_key:
                node_key = f"{node_key}-{len(node_by_key) + 1}"
            node_by_key[node_key] = node

        if len(node_by_key) < 2:
            raise RuntimeError("لم ينتج المقترح مفاهيم كافية لبناء شجرة.")

        # Normalize the AI relationships into a real tree: one root and one parent per node.
        node_keys = list(node_by_key)
        root_key = node_keys[0]
        tree_edges = []
        parent_by_target = {}
        for edge in payload["edges"]:
            source_key = str(edge.get("source"))
            target_key = str(edge.get("target"))
            if (
                source_key not in node_by_key
                or target_key not in node_by_key
                or source_key == target_key
                or target_key == root_key
                or target_key in parent_by_target
            ):
                continue
            cursor = source_key
            creates_cycle = False
            while cursor in parent_by_target:
                if cursor == target_key:
                    creates_cycle = True
                    break
                cursor = parent_by_target[cursor]
            if creates_cycle:
                continue
            tree_edges.append({
                "source": source_key,
                "target": target_key,
                "label": str(edge.get("label", ""))[:150] or "",
            })
            parent_by_target[target_key] = source_key

        # Attach any disconnected concepts to the root so the result is never a loose grid.
        for key in node_keys[1:]:
            if key not in parent_by_target:
                tree_edges.append({"source": root_key, "target": key, "label": "فرع"})
                parent_by_target[key] = root_key

        for edge in tree_edges:
            source = node_by_key[edge["source"]]
            target = node_by_key[edge["target"]]
            db.session.add(ConceptMapEdge(
                concept_map_id=row.id,
                source_node_id=source.id,
                target_node_id=target.id,
                label=edge["label"] or None,
            ))

        layout_concept_tree(node_by_key, tree_edges)
        db.session.commit()
        flash("أنشأ الذكاء الاصطناعي خريطة شجرية كمسودة. راجعيها قبل نشرها.", "success")
    except Exception as error:
        db.session.rollback()
        current_app.logger.exception("Concept map AI generation failed")
        flash(f"تعذر إنشاء الخريطة: {error}", "error")
    return redirect(url_for("platform_admin.concept_maps"))


@platform_admin_bp.post("/concept-maps/<int:row_id>/toggle")
def toggle_concept_map(row_id):
    row = owned_content_or_404(ConceptMap, row_id)
    row.publication_status = "draft" if row.publication_status == "published" else "published"
    db.session.commit()
    flash("تم تحديث حالة الخريطة.", "success")
    return redirect(url_for("platform_admin.concept_maps"))


@platform_admin_bp.post("/concept-maps/<int:row_id>/delete")
def delete_concept_map(row_id):
    row = owned_content_or_404(ConceptMap, row_id)
    relative_path = row.stored_path
    db.session.delete(row)
    db.session.commit()

    if relative_path:
        root = Path(current_app.config["UPLOAD_FOLDER"]).resolve()
        target = (root / relative_path).resolve()
        if target.is_relative_to(root) and target.is_file():
            try:
                target.unlink()
            except OSError:
                current_app.logger.warning("Could not remove concept map file %s", target)

    flash("تم حذف الخريطة المفاهيمية.", "success")
    return redirect(url_for("platform_admin.concept_maps"))


@platform_admin_bp.route("/characters", methods=["GET", "POST"])
def characters():
    form = HistoricalCharacterForm(); form.lesson_id.choices = lesson_choices()
    if form.validate_on_submit():
        db.session.add(HistoricalCharacter(lesson_id=form.lesson_id.data, created_by_id=current_user.id, name=form.name.data.strip(), period=(form.period.data or "").strip() or None, summary=form.summary.data.strip(), key_events=(form.key_events.data or "").strip() or None, clues=(form.clues.data or "").strip() or None, source_notes=(form.source_notes.data or "").strip() or None, publication_status=form.publication_status.data))
        db.session.commit(); flash("تم حفظ الشخصية التاريخية", "success")
        return redirect(url_for("platform_admin.characters"))
    query = HistoricalCharacter.query if current_user.role == "admin" else HistoricalCharacter.query.filter_by(created_by_id=current_user.id)
    return render_template("platform/characters.html", form=form, characters=query.order_by(HistoricalCharacter.created_at.desc()).all())


@platform_admin_bp.post("/characters/generate")
def generate_characters():
    lesson = db.session.get(Lesson, request.form.get("lesson_id", type=int))
    if lesson is None or lesson not in teacher_lessons():
        abort(404)
    try:
        drafts, citations = generate_character_drafts(lesson)
        created = 0
        for item in drafts:
            name = str(item.get("name", "")).strip()
            summary = str(item.get("summary", "")).strip()
            if not name or not summary or HistoricalCharacter.query.filter_by(lesson_id=lesson.id, name=name).first():
                continue
            clues = item.get("who_am_i_clues") or []
            notes = "\n".join(citations)
            db.session.add(HistoricalCharacter(
                lesson_id=lesson.id,
                created_by_id=current_user.id,
                name=name[:200],
                period=str(item.get("period", "")).strip()[:150] or None,
                summary=summary[:5000],
                key_events="\n".join(str(event) for event in (item.get("key_events") or []))[:5000] or None,
                source_notes=notes[:2000],
                clues="\n".join(str(clue) for clue in clues)[:3000] or None,
                publication_status="draft",
                is_ai_generated=True,
            ))
            created += 1
        db.session.commit()
        flash(f"تم حفظ {created} شخصية كمسودات للمراجعة.", "success")
    except Exception as error:
        db.session.rollback()
        current_app.logger.exception("Character AI generation failed")
        flash(f"تعذر استخراج الشخصيات: {error}", "error")
    return redirect(url_for("platform_admin.characters"))


@platform_admin_bp.route("/characters/<int:row_id>/edit", methods=["GET", "POST"])
def edit_character(row_id):
    row = owned_content_or_404(HistoricalCharacter, row_id)
    form = HistoricalCharacterForm(obj=row)
    form.lesson_id.choices = lesson_choices()
    if form.validate_on_submit():
        row.lesson_id = form.lesson_id.data
        row.name = form.name.data.strip()
        row.period = (form.period.data or "").strip() or None
        row.summary = form.summary.data.strip()
        row.key_events = (form.key_events.data or "").strip() or None
        row.clues = (form.clues.data or "").strip() or None
        row.source_notes = (form.source_notes.data or "").strip() or None
        row.publication_status = form.publication_status.data
        db.session.commit()
        flash("تمت مراجعة الشخصية وحفظها.", "success")
        return redirect(url_for("platform_admin.characters"))
    return render_template("platform/edit_character.html", form=form, character=row)


@platform_admin_bp.post("/characters/<int:row_id>/toggle")
def toggle_character(row_id):
    row = owned_content_or_404(HistoricalCharacter, row_id)
    row.publication_status = "draft" if row.publication_status == "published" else "published"
    db.session.commit()
    flash("تم تحديث ظهور الشخصية للطلاب.", "success")
    return redirect(url_for("platform_admin.characters"))


@platform_admin_bp.route("/settings", methods=["GET", "POST"])
def settings():
    row = PlatformSettings.get_or_create()
    form = PlatformSettingsForm(obj=row)

    if form.validate_on_submit():
        entered_whatsapp = (form.whatsapp_url.data or "").strip()
        whatsapp_value = None

        if entered_whatsapp:
            # Keep only secure, direct WhatsApp links from trusted WhatsApp domains.
            from urllib.parse import urlparse

            if entered_whatsapp.lower().startswith(("https://", "http://")):
                parsed = urlparse(entered_whatsapp)
                allowed_hosts = {
                    "wa.me",
                    "www.wa.me",
                    "api.whatsapp.com",
                    "chat.whatsapp.com",
                    "web.whatsapp.com",
                }
                if parsed.scheme == "https" and parsed.hostname and parsed.hostname.lower() in allowed_hosts:
                    whatsapp_value = entered_whatsapp
                else:
                    form.whatsapp_url.errors.append(
                        "أدخلي رقمًا مع مفتاح الدولة أو رابط واتساب مباشر يبدأ بـ https ومن نطاق واتساب."
                    )
                    return render_template("platform/settings.html", form=form)
            elif re.fullmatch(r"[\d\s()+.\-]+", entered_whatsapp):
                digits = re.sub(r"\D", "", entered_whatsapp)
                if digits.startswith("00"):
                    digits = digits[2:]
                elif digits.startswith("0"):
                    digits = "970" + digits[1:]

                if 8 <= len(digits) <= 15:
                    whatsapp_value = f"https://wa.me/{digits}"
                else:
                    form.whatsapp_url.errors.append(
                        "تأكدي من رقم واتساب ومفتاح الدولة؛ يجب أن يكون الرقم بين 8 و15 رقمًا."
                    )
                    return render_template("platform/settings.html", form=form)
            else:
                form.whatsapp_url.errors.append(
                    "أدخلي رقمًا فقط مع مفتاح الدولة أو رابط واتساب مباشر."
                )
                return render_template("platform/settings.html", form=form)

        try:
            row.platform_name = form.platform_name.data.strip()
            row.tagline = form.tagline.data.strip()
            row.whatsapp_url = whatsapp_value
            row.support_email = (form.support_email.data or "").strip() or None
            db.session.commit()
            flash("تم حفظ إعدادات المنصة وتحديث رابط واتساب.", "success")
            return redirect(url_for("platform_admin.settings"))
        except Exception:
            db.session.rollback()
            current_app.logger.exception("Platform settings save failed")
            flash("تعذر حفظ الإعدادات. لم يتم اعتماد التغييرات؛ حاولي مرة أخرى.", "error")

    return render_template("platform/settings.html", form=form)


@platform_admin_bp.get("/analytics")
def analytics():
    """Summarize assessment, practice, AI tutor, character, and points activity."""
    worksheet_query = Worksheet.query
    if current_user.role != "admin":
        worksheet_query = worksheet_query.filter_by(created_by_id=current_user.id)
    owned_worksheets = worksheet_query.all()
    worksheet_ids = [item.id for item in owned_worksheets]

    classroom_query = Classroom.query
    if current_user.role != "admin":
        classroom_query = classroom_query.filter_by(teacher_id=current_user.id)
    classrooms = classroom_query.all()
    student_ids = {
        enrollment.student_id
        for classroom in classrooms
        for enrollment in classroom.enrollments
    }

    if worksheet_ids:
        student_ids.update(
            student_id for (student_id,) in
            db.session.query(WorksheetAttempt.student_id)
            .filter(WorksheetAttempt.worksheet_id.in_(worksheet_ids))
            .distinct().all()
        )

    if current_user.role == "admin":
        student_ids.update(
            student_id for (student_id,) in
            db.session.query(User.id).filter_by(
                role="student", is_active_account=True
            ).all()
        )

    students = (
        User.query.filter(User.id.in_(student_ids), User.role == "student")
        .order_by(User.full_name).all()
        if student_ids else []
    )
    scoped_student_ids = [student.id for student in students]

    all_attempts = []
    if scoped_student_ids:
        all_attempts = WorksheetAttempt.query.filter(
            WorksheetAttempt.student_id.in_(scoped_student_ids),
        ).all()

    all_transactions = (
        XPTransaction.query.filter(
            XPTransaction.student_id.in_(scoped_student_ids)
        ).all()
        if scoped_student_ids else []
    )

    rows = []
    for student in students:
        student_all_attempts = [
            attempt for attempt in all_attempts
            if attempt.student_id == student.id
        ]
        student_attempts = [
            attempt for attempt in student_all_attempts
            if attempt.submitted_at is not None
        ]
        test_attempts = [
            attempt for attempt in student_attempts
            if attempt.worksheet and attempt.worksheet.creation_method == "test"
        ]
        worksheet_attempts = [
            attempt for attempt in student_attempts
            if not attempt.worksheet or attempt.worksheet.creation_method != "test"
        ]
        percentages = [attempt.percentage for attempt in student_attempts]
        mastery_rows = ConceptMastery.query.filter_by(student_id=student.id).all()
        challenge_sessions = ChallengeSession.query.filter_by(
            student_id=student.id
        ).all()
        challenges = [session for session in challenge_sessions if session.status == "completed"]
        tutor_interactions = TutorMessage.query.join(TutorConversation).filter(
            TutorConversation.student_id == student.id,
            TutorMessage.role == "student",
        ).count()
        character_attempts = CharacterAttempt.query.filter_by(
            student_id=student.id
        ).count()
        plan_items = LearningPlanItem.query.filter_by(student_id=student.id).all()
        plan_verifications = sum(1 for item in plan_items if item.verified_at is not None)
        student_transactions = [
            transaction for transaction in all_transactions
            if transaction.student_id == student.id
        ]
        resource_downloads = sum(1 for transaction in student_transactions if transaction.reason == "resource_downloaded")
        concept_map_views = sum(1 for transaction in student_transactions if transaction.reason == "concept_map_viewed")
        points_earned = sum(transaction.amount for transaction in student_transactions)
        average = round(sum(percentages) / len(percentages), 1) if percentages else 0.0
        mastery_average = (
            round(sum(item.mastery_score for item in mastery_rows) / len(mastery_rows), 1)
            if mastery_rows else 0.0
        )
        activity_score = (
            len(student_all_attempts) + len(challenge_sessions) + tutor_interactions
            + character_attempts + plan_verifications + resource_downloads
            + concept_map_views
        )
        rows.append({
            "student": student,
            "test_count": len(test_attempts),
            "worksheet_count": len(worksheet_attempts),
            "average": average,
            "mastery_average": mastery_average,
            "challenge_count": len(challenges),
            "tutor_interactions": tutor_interactions,
            "character_attempts": character_attempts,
            "plan_items": len(plan_items),
            "plan_verifications": plan_verifications,
            "resource_downloads": resource_downloads,
            "concept_map_views": concept_map_views,
            "points_earned": points_earned,
            "activity_score": activity_score,
            "needs_support": (
                (bool(percentages) and average < 50)
                or (bool(mastery_rows) and mastery_average < 50)
            ),
        })

    rows.sort(
        key=lambda item: (item["activity_score"], item["average"]),
        reverse=True,
    )
    submitted_attempts = [attempt for attempt in all_attempts if attempt.submitted_at is not None]
    class_percentages = [attempt.percentage or 0 for attempt in submitted_attempts]
    total_tutor_interactions = sum(row["tutor_interactions"] for row in rows)
    total_character_attempts = sum(row["character_attempts"] for row in rows)
    total_plan_verifications = sum(row["plan_verifications"] for row in rows)
    overview = {
        "students": len(students),
        "worksheets": sum(1 for item in owned_worksheets if item.creation_method != "test"),
        "tests": sum(1 for item in owned_worksheets if item.creation_method == "test"),
        "submissions": len(submitted_attempts),
        "average": round(sum(class_percentages) / len(class_percentages), 1)
        if class_percentages else 0.0,
        "challenge_completions": sum(
            ChallengeSession.query.filter_by(student_id=student.id, status="completed").count()
            for student in students
        ),
        "tutor_interactions": total_tutor_interactions,
        "character_attempts": total_character_attempts,
        "plan_verifications": total_plan_verifications,
        "resource_downloads": sum(row["resource_downloads"] for row in rows),
        "concept_map_views": sum(row["concept_map_views"] for row in rows),
        "points_awarded": sum(transaction.amount for transaction in all_transactions),
        "support_count": sum(1 for row in rows if row["needs_support"]),
    }
    return render_template("platform/analytics.html", rows=rows, overview=overview)


@platform_admin_bp.route("/honor-board", methods=["GET", "POST"])
def honor_board():
    classrooms = Classroom.query.filter_by(teacher_id=current_user.id, is_active=True).all()
    student_ids = {enrollment.student_id for classroom in classrooms for enrollment in classroom.enrollments}
    students = User.query.filter(User.id.in_(student_ids)).order_by(User.full_name).all() if student_ids else []
    if request.method == "POST":
        student_id = request.form.get("student_id", type=int)
        if student_id not in student_ids:
            abort(403)
        entry = HonorBoardEntry(teacher_id=current_user.id, student_id=student_id, category=request.form.get("category", "teacher_choice")[:50], note=(request.form.get("note") or "").strip()[:300] or None, period_label=(request.form.get("period_label") or "").strip()[:100] or None)
        db.session.add(entry); db.session.commit(); flash("تمت إضافة الطالبة إلى لوحة الشرف.", "success")
        return redirect(url_for("platform_admin.honor_board"))
    entries = HonorBoardEntry.query.filter_by(teacher_id=current_user.id, is_active=True).order_by(HonorBoardEntry.created_at.desc()).all()
    return render_template("platform/honor_board.html", students=students, entries=entries)


@platform_admin_bp.post("/honor-board/<int:entry_id>/remove")
def remove_honor_entry(entry_id):
    row = HonorBoardEntry.query.filter_by(id=entry_id, teacher_id=current_user.id).first_or_404()
    row.is_active = False
    db.session.commit(); flash("تمت إزالة البطاقة من لوحة الشرف.", "success")
    return redirect(url_for("platform_admin.honor_board"))
