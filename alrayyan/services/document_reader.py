import hashlib
from pathlib import Path

from docx import Document
import pymupdf


def calculate_checksum(file_path):
    """
    Create a unique SHA-256 fingerprint for a file.

    It helps us detect exact duplicate files.
    """
    file_path = Path(file_path)
    sha256 = hashlib.sha256()

    with file_path.open("rb") as file:
        for block in iter(lambda: file.read(8192), b""):
            sha256.update(block)

    return sha256.hexdigest()


def extract_pdf(file_path):
    """
    Extract text from a PDF while preserving page numbers.
    """
    document = pymupdf.open(str(file_path))
    pages = []

    try:
        for page_number, page in enumerate(
            document,
            start=1,
        ):
            text = page.get_text("text") or ""

            pages.append(
                {
                    "page_number": page_number,
                    "text": text.strip(),
                }
            )

        return {
            "document_type": "pdf",
            "page_count": document.page_count,
            "pages": pages,
        }

    finally:
        document.close()

def extract_docx(file_path):
    """
    Extract paragraphs and tables from a Word document.
    """
    document = Document(str(file_path))
    content = []

    for paragraph in document.paragraphs:
        text = paragraph.text.strip()

        if text:
            content.append(text)

    for table in document.tables:
        for row in table.rows:
            cells = [
                cell.text.strip()
                for cell in row.cells
                if cell.text.strip()
            ]

            if cells:
                content.append(" | ".join(cells))

    return {
        "document_type": "docx",
        "page_count": None,
        "text": "\n".join(content),
    }


def inspect_document(file_path):
    """
    Read a supported source and return its basic information.
    """
    file_path = Path(file_path)
    extension = file_path.suffix.lower()

    if extension == ".pdf":
        extracted = extract_pdf(file_path)
        text = "\n".join(
            page["text"]
            for page in extracted["pages"]
        )

    elif extension == ".docx":
        extracted = extract_docx(file_path)
        text = extracted["text"]

    else:
        raise ValueError(
            f"Unsupported file type: {extension}"
        )

    return {
        "filename": file_path.name,
        "path": str(file_path),
        "extension": extension,
        "checksum": calculate_checksum(file_path),
        "page_count": extracted["page_count"],
        "character_count": len(text),
        "text_preview": text[:300].replace("\n", " "),
    }